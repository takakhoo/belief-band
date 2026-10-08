"""Load every raw file under data/raw and print its date range and row count.

Needs pandas, openpyxl, xlrd (for the two .xls files). Run from anywhere:
    python3 scripts/check_raw_data.py
"""
import re
import sys
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import openpyxl.reader.excel as _rx
import pandas as pd

# Some Philly Fed workbooks carry a malformed docProps/core.xml that openpyxl rejects.
_rx.ExcelReader.read_properties = lambda self: None

RAW = Path(__file__).resolve().parent.parent / "data" / "raw"
OUT = []


def emit(path, sheet, first, last, n, note=""):
    rel = str(Path(path).relative_to(RAW))
    OUT.append((rel, sheet, str(first), str(last), n, note))


def parse_date(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v
    if hasattr(v, "year") and hasattr(v, "month"):
        return pd.Timestamp(v)
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return pd.to_datetime(s, format=fmt)
        except ValueError:
            pass
    return None


def has_numeric(row):
    for c in row:
        if isinstance(c, (int, float)) and not (isinstance(c, float) and np.isnan(c)):
            return True
        if isinstance(c, str):
            try:
                float(c)
                return True
            except ValueError:
                pass
    return False


# ---------- FRED ----------
def check_fred():
    for f in sorted((RAW / "fred").glob("*.csv")):
        df = pd.read_csv(f, na_values=["."])
        col = df.columns[1]
        ok = df[df[col].notna()]
        note = f"{len(df) - len(ok)} blank obs" if len(df) != len(ok) else ""
        emit(f, col, ok.iloc[0, 0], ok.iloc[-1, 0], len(ok), note)


# ---------- Ken French ----------
def check_french():
    for f in sorted((RAW / "french").glob("*.csv")):
        lines = f.read_text(encoding="latin1").splitlines()
        blocks, cur = [], []
        for ln in lines:
            m = re.match(r"^\s*(\d{4,8})\s*,", ln)
            if m:
                cur.append(m.group(1))
            elif cur:
                blocks.append(cur)
                cur = []
        if cur:
            blocks.append(cur)
        main = blocks[0]
        freq = {4: "annual", 6: "monthly", 8: "daily"}[len(main[0])]
        n_ann = sum(1 for b in blocks if len(b[0]) == 4)
        emit(f, f"block1 ({freq})", main[0], main[-1], len(main),
             f"{len(blocks)} data blocks ({n_ann} annual)")


# ---------- generic xlsx table: first column dates ----------
def xlsx_date_table(path, sheet, header_row=None, label=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    dates, n = [], 0
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if header_row is not None and i <= header_row:
            continue
        if not r:
            continue
        d = parse_date(r[0])
        if d is None:
            continue
        if has_numeric(r[1:]):
            dates.append(d)
            n += 1
    wb.close()
    if dates:
        emit(path, label or sheet, min(dates).date(), max(dates).date(), n)
    else:
        emit(path, label or sheet, "-", "-", 0, "no dated numeric rows")


# ---------- Goyal ----------
def check_goyal():
    d = RAW / "goyal"
    for f in sorted(d.glob("*.xlsx")):
        x = pd.read_excel(f, sheet_name=None, na_values=["NaN"])
        for name, df in x.items():
            if name == "ReadMe":
                continue
            key = df.columns[0]
            df = df[pd.to_numeric(df[key], errors="coerce").notna()]
            ret_col = "CRSP_SPvw" if "CRSP_SPvw" in df.columns else "ret"
            r = df[df[ret_col].notna()]
            emit(f, name, int(df[key].iloc[0]), int(df[key].iloc[-1]), len(df),
                 f"{ret_col} non-missing {int(r[key].iloc[0])}-{int(r[key].iloc[-1])}")
    z = d / "GWZ2024_data_2025_csv.zip"
    with zipfile.ZipFile(z) as zf:
        names = zf.namelist()
        with zf.open("gip_A.csv") as fh:
            g = pd.read_csv(fh, header=None, na_values=["NaN"])
    emit(z, f"{len(names)} csv (unextracted)", int(g.iloc[1, 0]), int(g.iloc[-1, 0]), len(g) - 1,
         "sample member gip_A.csv: rows=obs year, cols=vintage/forecast-origin")


# ---------- Shiller ----------
def check_shiller():
    for f in sorted((RAW / "shiller").glob("*.xls")):
        df = pd.read_excel(f, sheet_name="Data", header=None, skiprows=8, engine="xlrd")
        df = df[pd.to_numeric(df[0], errors="coerce").notna() & pd.to_numeric(df[1], errors="coerce").notna()]
        # Date is yyyy.mm stored as float, so October shows as yyyy.1
        def fmt(v):
            y = int(v)
            m = int(round((v - y) * 100))
            return f"{y}-{m:02d}"
        emit(f, "Data", fmt(df[0].iloc[0]), fmt(df[0].iloc[-1]), len(df))


# ---------- Philly Fed RTDSM ----------
def check_rtdsm():
    for f in sorted((RAW / "philfed_rtdsm").glob("*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            if rows[0][0] == "DATE":
                hdr = [h for h in rows[0][1:] if h]
                obs = [r[0] for r in rows[1:] if r and r[0]]
                emit(f, ws.title, obs[0], obs[-1], len(obs),
                     f"{len(hdr)} vintages {hdr[0]}..{hdr[-1]}")
            else:
                obs = [r[0] for r in rows if r and isinstance(r[0], str) and re.match(r"^\d{4}:(Q?\d)", r[0])]
                if obs:
                    emit(f, ws.title, obs[0], obs[-1], len(obs), "first/second/third release table")
        wb.close()


# ---------- JST ----------
def check_jst():
    d = RAW / "jst"
    x = pd.read_excel(d / "JSTdatasetR6.xlsx")
    emit(d / "JSTdatasetR6.xlsx", "Sheet1", x.year.min(), x.year.max(), len(x),
         f"{x.country.nunique()} countries; eq_tr non-missing to {int(x.loc[x.eq_tr.notna(), 'year'].max())}")
    s = pd.read_stata(d / "JSTdatasetR6.dta")
    emit(d / "JSTdatasetR6.dta", "stata", int(s.year.min()), int(s.year.max()), len(s),
         f"{s.shape[1]} vars")


# ---------- AQR ----------
def check_aqr():
    d = RAW / "aqr"
    for f in sorted(d.glob("*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True)
        sheets = [s for s in wb.sheetnames
                  if s not in ("Definition", "Definitions", "Data Sources", "Disclosures",
                               "Sources and Definitions", "--> Additional Global Factors",
                               "Series Information")]
        wb.close()
        for s in sheets:
            xlsx_date_table(f, s)


# ---------- Treasury / Fed / Damodaran ----------
def check_treasury():
    d = RAW / "treasury"
    g = pd.read_csv(d / "feds200628.csv", skiprows=9)
    ok = g[g["SVENY10"].notna()]
    emit(d / "feds200628.csv", "GSW daily", g.Date.iloc[0], g.Date.iloc[-1], len(g),
         f"SVENY10 non-missing {ok.Date.iloc[0]}..{ok.Date.iloc[-1]}; SVENY30 from "
         f"{g.loc[g.SVENY30.notna(), 'Date'].iloc[0]}")
    h = pd.read_excel(d / "histretSP.xls", sheet_name="Returns by year", header=None, engine="xlrd")
    yrs = pd.to_numeric(h[0], errors="coerce")
    yrs = yrs[(yrs > 1900) & (yrs < 2100)]
    emit(d / "histretSP.xls", "Returns by year", int(yrs.min()), int(yrs.max()), len(yrs), "annual")


# ---------- Schwert ----------
def check_schwert():
    d = RAW / "schwert"
    dd = pd.read_csv(d / "STKDATD.DAT", sep=r"\s+", header=None)
    emit(d / "STKDATD.DAT", "daily", dd[0].iloc[0], dd[0].iloc[-1], len(dd))
    mm = pd.read_csv(d / "STKDATM.DAT", sep=r"\s+", header=None, nrows=1488)
    emit(d / "STKDATM.DAT", "monthly", mm[0].iloc[0], mm[0].iloc[-1], len(mm))


def main():
    for fn in (check_fred, check_french, check_goyal, check_shiller, check_rtdsm,
               check_jst, check_aqr, check_treasury, check_schwert):
        try:
            fn()
        except Exception as e:  # keep going so one bad file doesn't hide the rest
            OUT.append((fn.__name__, "ERROR", "", "", 0, repr(e)[:200]))
    print("| file | sheet/series | first | last | rows | notes |")
    print("|---|---|---|---|---|---|")
    for r in OUT:
        print("| " + " | ".join(str(x) for x in r) + " |")


if __name__ == "__main__":
    sys.exit(main())
